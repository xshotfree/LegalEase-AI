from __future__ import annotations

import io
import re

from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF

from PIL import Image


def sanitize_text(text: str) -> str:

    replacements = {

        "\u2018": "'",

        "\u2019": "'",

        "\u201c": '"',

        "\u201d": '"',

        "\u2013": "-",

        "\u2014": "-",

        "\u00a0": " ",

        "\u2022": "-",

        "\u200b": "",
    }

    for source, target in replacements.items():

        text = text.replace(
            source,
            target
        )

    return text.strip()


def _split_lines(text: str) -> list[str]:

    return [
        line.strip()
        for line
        in sanitize_text(text).splitlines()
    ]


def _title_from_content(
    text: str,
    fallback: str
) -> str:

    for line in _split_lines(text):

        if line:

            return line[:160]

    return fallback


def make_txt(text: str) -> bytes:

    return sanitize_text(text).encode(
        "utf-8"
    )


def make_docx(
    text: str,
    doc_type: str,
    terms: str = "",
    brand_name: str = "LegalEase",
    footer_text: str = "",
    logo_bytes: Optional[bytes] = None,
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)

    section.bottom_margin = Inches(0.7)

    section.left_margin = Inches(0.85)

    section.right_margin = Inches(0.85)

    styles = document.styles

    styles["Normal"].font.name = (
        "Times New Roman"
    )

    styles["Normal"].font.size = Pt(11)

    # Logo
    if logo_bytes:

        try:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run()

            run.add_picture(
                io.BytesIO(logo_bytes),
                width=Inches(1.35)
            )

        except Exception:

            pass

    # Title
    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        _title_from_content(
            text,
            doc_type
        )
    )

    run.bold = True

    run.font.name = "Times New Roman"

    run.font.size = Pt(16)

    # Brand
    brand = document.add_paragraph()

    brand.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    brand_run = brand.add_run(
        brand_name
    )

    brand_run.italic = True

    brand_run.font.size = Pt(9)

    # Main document
    for line in _split_lines(text):

        if not line:

            document.add_paragraph("")

            continue

        cleaned = line.strip()

        is_heading = (

            len(cleaned) <= 90

            and (

                cleaned.isupper()

                or re.match(
                    r"^\d+[\.\)]\s+[A-Z]",
                    cleaned
                )

                or cleaned.endswith(":")
            )
        )

        paragraph = document.add_paragraph()

        if is_heading:

            run = paragraph.add_run(
                cleaned
            )

            run.bold = True

            run.font.name = (
                "Times New Roman"
            )

            run.font.size = Pt(12)

        else:

            paragraph.paragraph_format.space_after = Pt(5)

            paragraph.add_run(
                cleaned
            )

    # Terms table
    term_items = [

        item.strip()

        for item in terms.replace(
            "\n",
            ";"
        ).split(";")

        if item.strip()
    ]

    if term_items:

        document.add_page_break()

        heading = document.add_paragraph()

        heading_run = heading.add_run(
            "Terms Table"
        )

        heading_run.bold = True

        heading_run.font.size = Pt(13)

        table = document.add_table(
            rows=1,
            cols=2
        )

        table.style = "Table Grid"

        table.rows[0].cells[0].text = "No."

        table.rows[0].cells[1].text = (
            "Term / Condition"
        )

        for index, item in enumerate(
            term_items,
            start=1
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)

            cells[1].text = item

    # Footer
    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(

        footer_text

        or (
            "Generated with LegalEase. "
            "Review by a qualified legal "
            "professional is recommended."
        )
    )

    footer_run.font.size = Pt(8)

    output = io.BytesIO()

    document.save(output)

    return output.getvalue()


def _safe_pdf_text(text: str) -> str:

    text = sanitize_text(text)

    return (
        text
        .encode(
            "latin-1",
            errors="replace"
        )
        .decode("latin-1")
    )


class LegalEasePDF(FPDF):

    def __init__(
        self,
        brand_name: str,
        footer_text: str,
        logo_path: Optional[str] = None,
    ):

        super().__init__()

        self.brand_name = brand_name

        self.footer_text = footer_text

        self.logo_path = logo_path

    def header(self):

        if (
            self.logo_path
            and Path(
                self.logo_path
            ).exists()
        ):

            try:

                self.image(
                    self.logo_path,
                    x=10,
                    y=8,
                    w=18
                )

                self.set_xy(
                    32,
                    10
                )

            except Exception:

                self.set_xy(
                    10,
                    10
                )

        else:

            self.set_xy(
                10,
                10
            )

        self.set_font(
            "Helvetica",
            "B",
            10
        )

        self.cell(
            0,
            8,
            _safe_pdf_text(
                self.brand_name
            ),
            align="L"
        )

        self.ln(12)

    def footer(self):

        self.set_y(-18)

        self.set_font(
            "Helvetica",
            "",
            7
        )

        self.cell(
            0,
            8,
            _safe_pdf_text(
                self.footer_text
            ),
            align="C"
        )


def make_pdf(
    text: str,
    doc_type: str,
    brand_name: str = "LegalEase",
    footer_text: str = "",
    logo_bytes: Optional[bytes] = None,
) -> bytes:

    logo_path = None

    temp_path = None

    if logo_bytes:

        try:

            image = Image.open(
                io.BytesIO(logo_bytes)
            )

            temp_path = (
                Path(__file__).resolve().parent
                / "_temp_logo.png"
            )

            image.convert(
                "RGB"
            ).save(
                temp_path,
                format="PNG"
            )

            logo_path = str(
                temp_path
            )

        except Exception:

            logo_path = None

    try:

        pdf = LegalEasePDF(

            brand_name=brand_name,

            footer_text=(
                footer_text

                or (
                    "Generated with LegalEase. "
                    "Review by a qualified legal "
                    "professional is recommended."
                )
            ),

            logo_path=logo_path,
        )

        pdf.set_auto_page_break(
            auto=True,
            margin=22
        )

        pdf.add_page()

        title = _title_from_content(
            text,
            doc_type
        )

        pdf.set_font(
            "Helvetica",
            "B",
            15
        )

        pdf.multi_cell(
            0,
            8,
            _safe_pdf_text(title),
            align="C"
        )

        pdf.ln(4)

        for line in _split_lines(text):

            if not line:

                pdf.ln(3)

                continue

            is_heading = (

                len(line) <= 90

                and (

                    line.isupper()

                    or re.match(
                        r"^\d+[\.\)]\s+[A-Z]",
                        line
                    )

                    or line.endswith(":")
                )
            )

            pdf.set_font(
                "Helvetica",
                "B" if is_heading else "",
                11 if is_heading else 10
            )

            pdf.multi_cell(
                0,
                6,
                _safe_pdf_text(line)
            )

            pdf.ln(1)

        return bytes(
            pdf.output()
        )

    finally:

        if (
            temp_path
            and temp_path.exists()
        ):

            try:

                temp_path.unlink()

            except OSError:

                pass


def format_html_preview(
    text: str
) -> str:

    import html

    escaped = html.escape(
        sanitize_text(text)
    )

    return escaped.replace(
        "\n",
        "<br>"
    )